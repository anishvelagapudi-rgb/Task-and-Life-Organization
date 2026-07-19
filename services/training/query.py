import csv
import io
import json
from datetime import datetime

# Shared read/aggregation logic over training_extractions -- used by the
# dashboard route, the single-metric chart route, and the AI's
# query_training_data/export_training_data/graph_training_metric tools, so
# these three surfaces can't drift into three slightly-different answers to
# the same question. Every function here is a pure function of (db, filters)
# -- no Flask/AI-provider imports -- so it's equally callable from a route or
# a tool executor.

CHARTED_METRIC_TYPES = ("weight", "run", "workout_set")


def _date_filtered(db, metric_type: str, date_from: str | None, date_to: str | None):
    sql = "SELECT entry_date, data FROM training_extractions WHERE metric_type = ? AND superseded = 0"
    params: list = [metric_type]
    if date_from:
        sql += " AND entry_date >= ?"
        params.append(date_from)
    if date_to:
        sql += " AND entry_date <= ?"
        params.append(date_to)
    sql += " ORDER BY entry_date"
    return [dict(r) for r in db.execute(sql, params).fetchall()]


def weight_trend(db, date_from: str | None = None, date_to: str | None = None) -> list[dict]:
    points = []
    for r in _date_filtered(db, "weight", date_from, date_to):
        try:
            v = json.loads(r["data"]).get("value_lbs")
        except (json.JSONDecodeError, AttributeError):
            v = None
        if isinstance(v, (int, float)):
            points.append({"date": r["entry_date"], "value_lbs": v})
    return points


def weekly_mileage(db, date_from: str | None = None, date_to: str | None = None) -> list[dict]:
    totals: dict[str, float] = {}
    for r in _date_filtered(db, "run", date_from, date_to):
        try:
            dist = json.loads(r["data"]).get("distance_mi")
        except (json.JSONDecodeError, AttributeError):
            dist = None
        if not isinstance(dist, (int, float)):
            continue
        year, week, _ = datetime.fromisoformat(r["entry_date"]).isocalendar()
        label = f"{year}-W{week:02d}"
        totals[label] = totals.get(label, 0) + dist
    return [{"week": k, "miles": round(v, 2)} for k, v in sorted(totals.items())]


def one_rm_by_exercise(db, date_from: str | None = None, date_to: str | None = None) -> dict[str, list[dict]]:
    by_exercise: dict[str, list[dict]] = {}
    for r in _date_filtered(db, "workout_set", date_from, date_to):
        try:
            d = json.loads(r["data"])
        except (json.JSONDecodeError, TypeError):
            continue
        exercise, reps, weight_lbs = d.get("exercise"), d.get("reps"), d.get("weight_lbs")
        if not exercise or not isinstance(reps, (int, float)) or not isinstance(weight_lbs, (int, float)):
            continue
        one_rm = round(weight_lbs * (1 + reps / 30.0), 1)  # Epley formula
        key = str(exercise).strip().lower()
        by_exercise.setdefault(key, []).append({"date": r["entry_date"], "one_rm": one_rm})
    return by_exercise


# Which JSON field to plot for metric types without bespoke chart logic above.
_GENERIC_FIELD_BY_TYPE = {
    "sleep": "hours",
    "resting_hr": "bpm",
    "steps": "count",
}


def metric_series(db, metric_type: str, date_from: str | None = None, date_to: str | None = None) -> list[dict]:
    """Generic {date, value} extractor for any metric_type not covered by the
    bespoke functions above -- used by the chart route/graph_training_metric
    tool so every metric_type is graphable, not just the three Phase 1 built
    dedicated logic for."""
    field = _GENERIC_FIELD_BY_TYPE.get(metric_type)
    points = []
    for r in _date_filtered(db, metric_type, date_from, date_to):
        try:
            d = json.loads(r["data"])
        except (json.JSONDecodeError, TypeError):
            continue
        value = d.get(field) if field else next((v for v in d.values() if isinstance(v, (int, float))), None)
        if isinstance(value, (int, float)):
            points.append({"date": r["entry_date"], "value": value})
    return points


def other_extractions(db, exclude_types=CHARTED_METRIC_TYPES, limit: int = 50) -> list[dict]:
    placeholders = ", ".join("?" * len(exclude_types))
    rows = [dict(r) for r in db.execute(
        f"SELECT source_entry_id, entry_date, metric_type, data, confidence, extracted_at "
        f"FROM training_extractions WHERE metric_type NOT IN ({placeholders}) AND superseded = 0 "
        f"ORDER BY extracted_at DESC LIMIT ?",
        [*exclude_types, limit],
    ).fetchall()]
    for r in rows:
        try:
            r["data"] = json.loads(r["data"])
        except (json.JSONDecodeError, TypeError):
            r["data"] = {}
    return rows


def query_rows(
    db,
    metric_types: list[str] | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    keyword: str | None = None,
    limit: int = 300,
) -> list[dict]:
    """Filtered training_extractions rows, joined back to the source entry's raw
    content ("linking back to original journal entries" per the spec). No
    numeric-threshold parsing (e.g. "long runs") -- callers (the AI tool) get the
    real rows for the window and reason over them, the same way service.py
    already dumps the full task list into context rather than pre-filtering it."""
    sql = """
        SELECT x.source_entry_id, x.entry_date, x.metric_type, x.data, x.confidence,
               x.extracted_at, e.content AS entry_content
        FROM training_extractions x
        LEFT JOIN training_entries e ON e.id = x.source_entry_id
        WHERE x.superseded = 0
    """
    params: list = []
    if metric_types:
        sql += f" AND x.metric_type IN ({', '.join('?' * len(metric_types))})"
        params.extend(metric_types)
    if date_from:
        sql += " AND x.entry_date >= ?"
        params.append(date_from)
    if date_to:
        sql += " AND x.entry_date <= ?"
        params.append(date_to)
    if keyword:
        sql += " AND (x.data ILIKE ? OR e.content ILIKE ?)"
        like = f"%{keyword}%"
        params.extend([like, like])
    sql += " ORDER BY x.entry_date LIMIT ?"
    params.append(limit)

    rows = [dict(r) for r in db.execute(sql, params).fetchall()]
    for r in rows:
        try:
            r["data"] = json.loads(r["data"])
        except (json.JSONDecodeError, TypeError):
            r["data"] = {}
    return rows


def export_rows(rows: list[dict], fmt: str) -> tuple[bytes, str, str]:
    """Returns (bytes, content_type, file_extension)."""
    columns = ["entry_date", "metric_type", "data", "confidence", "entry_content"]
    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "training_data"
        ws.append(columns)
        for r in rows:
            ws.append([
                r.get("entry_date"), r.get("metric_type"),
                json.dumps(r.get("data")), r.get("confidence"), r.get("entry_content"),
            ])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(columns)
    for r in rows:
        writer.writerow([
            r.get("entry_date"), r.get("metric_type"),
            json.dumps(r.get("data")), r.get("confidence"), r.get("entry_content"),
        ])
    return buf.getvalue().encode("utf-8"), "text/csv", "csv"
